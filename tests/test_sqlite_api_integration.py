from __future__ import annotations

from pathlib import Path
from io import BytesIO

from docx import Document
import fitz

from fastapi.testclient import TestClient

from backend.app.main import create_app
from backend.app.settings import Settings


def test_asset_api_survives_application_recreation(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path)
    first = TestClient(create_app(settings))
    response = first.post(
        "/api/v1/assets",
        files={"file": ("paper.md", b"# Title\n\nEvidence\n", "text/markdown")},
    )
    assert response.status_code == 201
    asset_id = response.json()["id"]

    second = TestClient(create_app(settings))
    fetched = second.get(f"/api/v1/assets/{asset_id}")
    assert fetched.status_code == 200
    assert fetched.json()["sha256"] == response.json()["sha256"]
    assert fetched.json()["blocks"][0]["asset_id"] == asset_id

    duplicate = first.post(
        "/api/v1/assets",
        files={"file": ("renamed.md", b"# Title\n\nEvidence\n", "text/markdown")},
    )
    assert duplicate.status_code == 201
    assert duplicate.json()["id"] == asset_id


def test_rubric_api_persists_draft_and_human_publication(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path)
    first = TestClient(create_app(settings))
    source = first.post(
        "/api/v1/assets",
        files={"file": ("rubric.md", b"# Rubric\n", "text/markdown")},
    )
    assert source.status_code == 201
    imported = first.post(
        "/api/v1/rubrics/import",
        json={
            "format": "json",
            "source_asset_id": source.json()["id"],
            "source_hash": source.json()["sha256"],
            "content": {
                "name": "官方量表",
                "version": "v1",
                "criteria": [{"id": "model", "label": "模型", "weight": 1}],
            },
        },
    )
    assert imported.status_code == 201
    rubric_id = imported.json()["id"]

    published = first.post(
        f"/api/v1/rubrics/{rubric_id}/publish",
        json={"confirmed_by": "human", "decision": "publish", "rationale": "已核对"},
    )
    assert published.status_code == 200
    assert published.json()["rubric"]["status"] == "published"

    second = TestClient(create_app(settings))
    persisted = second.get(f"/api/v1/rubrics/{rubric_id}")
    assert persisted.status_code == 200
    assert persisted.json()["status"] == "published"
    assert persisted.json()["decision_ledger_id"] == published.json()["rubric"]["decision_ledger_id"]


def test_docx_rubric_upload_is_parsed_through_asset_inventory(tmp_path: Path) -> None:
    document = Document()
    document.add_heading("DOCX 评分", level=1)
    table = document.add_table(rows=2, cols=6)
    headers = ["id", "label", "description", "weight", "scale_min", "scale_max"]
    values = ["model", "模型", "质量", "1", "0", "100"]
    for index, value in enumerate(headers):
        table.cell(0, index).text = value
    for index, value in enumerate(values):
        table.cell(1, index).text = value
    stream = BytesIO()
    document.save(stream)

    client = TestClient(create_app(Settings(data_dir=tmp_path)))
    response = client.post(
        "/api/v1/rubrics/import",
        files={"file": ("rubric.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 201
    assert response.json()["name"] == "DOCX 评分"
    assert response.json()["criteria"][0]["id"] == "model"


def test_pdf_rubric_upload_is_parsed_through_asset_inventory(tmp_path: Path) -> None:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "id | label | description | weight | scale_min | scale_max")
    page.insert_text((72, 90), "---|---|---|---|---|---")
    page.insert_text((72, 108), "model | Model | Quality | 1 | 0 | 100")
    stream = BytesIO(document.tobytes())
    document.close()

    client = TestClient(create_app(Settings(data_dir=tmp_path)))
    response = client.post(
        "/api/v1/rubrics/import",
        files={"file": ("rubric.pdf", stream.getvalue(), "application/pdf")},
    )
    assert response.status_code == 201
    assert response.json()["criteria"][0]["id"] == "model"


def test_rubric_source_asset_and_hash_are_verified(tmp_path: Path) -> None:
    client = TestClient(create_app(Settings(data_dir=tmp_path)))
    source = client.post("/api/v1/assets", files={"file": ("source.md", b"# Source\n", "text/markdown")})
    assert source.status_code == 201
    base = {
        "format": "json",
        "source_asset_id": source.json()["id"],
        "source_hash": source.json()["sha256"],
        "content": {"name": "Verified", "version": "v1", "criteria": [{"id": "m", "label": "模型", "weight": 1}]},
    }
    assert client.post("/api/v1/rubrics/import", json=base).status_code == 201
    bad_id = dict(base, source_asset_id="missing")
    missing = client.post("/api/v1/rubrics/import", json=bad_id)
    assert missing.status_code == 404
    assert missing.json()["error"]["code"] == "source_asset_not_found"
    bad_hash = dict(base, source_hash="bad")
    mismatched = client.post("/api/v1/rubrics/import", json=bad_hash)
    assert mismatched.status_code == 409
    assert mismatched.json()["error"]["code"] == "source_hash_mismatch"


def test_failed_asset_is_persisted_for_audit(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path)
    client = TestClient(create_app(settings))
    response = client.post("/api/v1/assets", files={"file": ("empty.md", b"", "text/markdown")})
    assert response.status_code == 422
    asset_id = response.json()["error"]["details"]["id"]
    fetched = client.get(f"/api/v1/assets/{asset_id}")
    assert fetched.status_code == 200
    assert fetched.json()["status"] == "failed"


def test_review_job_binds_published_inputs_and_persists_state(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path)
    client = TestClient(create_app(settings))
    paper = client.post("/api/v1/assets", files={"file": ("paper.md", b"# Paper\n\nEvidence\n", "text/markdown")})
    source = client.post("/api/v1/assets", files={"file": ("rubric.md", b"# Rubric\n", "text/markdown")})
    assert paper.status_code == source.status_code == 201
    imported = client.post(
        "/api/v1/rubrics/import",
        json={
            "format": "json",
            "source_asset_id": source.json()["id"],
            "source_hash": source.json()["sha256"],
            "content": {"name": "Official", "version": "v1", "criteria": [{"id": "m", "label": "模型", "weight": 1}]},
        },
    )
    assert imported.status_code == 201
    rubric_id = imported.json()["id"]
    assert client.post(f"/api/v1/rubrics/{rubric_id}/publish", json={"confirmed_by": "human", "rationale": "已确认"}).status_code == 200
    created = client.post("/api/v1/reviews", json={"asset_id": paper.json()["id"], "rubric_profile_id": rubric_id})
    assert created.status_code == 201
    review_id = created.json()["id"]
    transitioned = client.post(f"/api/v1/reviews/{review_id}/transition", json={"target": "INGESTING", "reason": "开始解析"})
    assert transitioned.status_code == 200
    assert transitioned.json()["state"] == "INGESTING"
    stage = client.post(
        f"/api/v1/reviews/{review_id}/stages/problem_identification",
        json={
            "stage_id": "problem_identification",
            "raw_score_0_10": 8,
            "rationale": "已定位题目证据",
            "evidence": [{"id": "ev-stage", "kind": "document", "asset_id": paper.json()["id"], "block_id": paper.json()["blocks"][0]["block_id"], "line_start": 1, "evidence_level": "VERIFIED"}],
        },
    )
    assert stage.status_code == 200
    result = client.put(f"/api/v1/reviews/{review_id}/result", json={"evidence": [{"id": "ev-result", "kind": "document", "asset_id": paper.json()["id"], "block_id": paper.json()["blocks"][0]["block_id"], "line_start": 1, "evidence_level": "VERIFIED"}]})
    assert result.status_code == 200
    assert client.get(f"/api/v1/reviews/{review_id}/result").status_code == 200
    restarted = TestClient(create_app(settings))
    persisted = restarted.get(f"/api/v1/reviews/{review_id}")
    assert persisted.status_code == 200
    assert persisted.json()["state_snapshot"]["events"][0]["to_state"] == "INGESTING"
