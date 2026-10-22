from __future__ import annotations

from backend.app.domain.assets import DocumentBlock


def parse_docx(path, asset_id: str) -> tuple[DocumentBlock, ...]:
    """Extract paragraphs and tables from a DOCX with stable document order."""

    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - exercised by service error path
        raise RuntimeError("missing_dependency:python-docx") from exc

    document = Document(str(path))
    blocks: list[DocumentBlock] = []
    order = 0
    paragraph_number = 0
    table_number = 0
    paragraphs_by_xml = {id(paragraph._p): paragraph for paragraph in document.paragraphs}
    tables_by_xml = {id(table._tbl): table for table in document.tables}
    # Walk the body XML so paragraph/table blocks retain their actual order.
    for child in document.element.body.iterchildren():
        paragraph = paragraphs_by_xml.get(id(child))
        table = tables_by_xml.get(id(child))
        if paragraph is not None:
            paragraph_number += 1
            text = paragraph.text.strip()
            if not text:
                continue
            order += 1
            blocks.append(
                DocumentBlock(
                    block_id=f"{asset_id}:block-{order:04d}-paragraph-{paragraph_number}",
                    asset_id=asset_id,
                    block_type="heading" if paragraph.style and paragraph.style.name.lower().startswith("heading") else "paragraph",
                    text=text,
                    paragraph=paragraph_number,
                    block_order=order,
                )
            )
            continue
        if table is None:
            continue
        table_number += 1
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        text = "\n".join("\t".join(row) for row in rows).strip()
        if not text:
            continue
        order += 1
        cells = [
            {"row": row_index, "column": column_index, "text": value}
            for row_index, row in enumerate(rows, start=1)
            for column_index, value in enumerate(row, start=1)
        ]
        blocks.append(
            DocumentBlock(
                block_id=f"{asset_id}:block-{order:04d}-table-{table_number}",
                asset_id=asset_id,
                block_type="table",
                text=text,
                block_order=order,
                metadata={
                    "table_number": table_number,
                    "rows": len(rows),
                    "columns": max((len(row) for row in rows), default=0),
                    "cells": cells,
                },
            )
        )
    return tuple(blocks)
